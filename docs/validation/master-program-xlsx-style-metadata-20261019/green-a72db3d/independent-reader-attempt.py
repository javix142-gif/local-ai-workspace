import json, sys
from openpyxl import load_workbook
path=sys.argv[1]
wb=load_workbook(path,data_only=False,read_only=False)
ws=wb['Ventas']
result={
 'reader':'openpyxl',
 'version':__import__('openpyxl').__version__,
 'sheetNames':wb.sheetnames,
 'A1StyleId':ws['A1'].style_id,
 'A1NumberFormat':ws['A1'].number_format,
 'A1ReaderValue':str(ws['A1'].value),
 'B1StyleId':ws['B1'].style_id,
 'B1NumberFormat':ws['B1'].number_format,
 'B1ReaderValue':str(ws['B1'].value),
 'D1SharedString':ws['D1'].value,
 'E1InlineString':ws['E1'].value,
 'F1Formula':ws['F1'].value,
 'F1StyleId':ws['F1'].style_id,
 'zipEntries':len(wb._archive.namelist()) if hasattr(wb,'_archive') else None,
}
print(json.dumps(result,ensure_ascii=False,indent=2))
assert result['sheetNames']==['Ventas']
assert result['A1StyleId']==1 and result['A1NumberFormat']=='mm-dd-yy'
assert result['B1StyleId']==2 and result['B1NumberFormat']=='yyyy-mm-dd'
assert result['D1SharedString']=='Shared label' and result['E1InlineString']=='Inline label'
assert result['F1Formula']=='=A1+1'
