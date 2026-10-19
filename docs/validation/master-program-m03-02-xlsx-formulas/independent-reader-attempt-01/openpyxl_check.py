import hashlib, json, sys, zipfile
from pathlib import Path
import openpyxl
root=Path(__file__).parent/'reader-output'
def inspect(name, expected_formula, expected_cached):
    path=root/name
    with zipfile.ZipFile(path) as archive:
        bad=archive.testzip()
        if bad is not None: raise AssertionError(f'{name}: corrupt ZIP member {bad}')
    formulas=openpyxl.load_workbook(path,data_only=False,read_only=False)
    values=openpyxl.load_workbook(path,data_only=True,read_only=False)
    sheet=formulas['Hoja1']
    cached=values['Hoja1']['A3'].value
    formula=sheet['A3'].value
    assert formula==expected_formula,(name,formula,expected_formula)
    assert cached==expected_cached,(name,cached,expected_cached)
    assert sheet['A1'].value==1 and sheet['A2'].value==5
    calc=formulas.calculation.fullCalcOnLoad
    if name=='formula-edited.xlsx': assert calc is True,calc
    return {'name':name,'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'zipCrc':'PASS','sheetNames':formulas.sheetnames,'A1':sheet['A1'].value,'A2':sheet['A2'].value,'A3Formula':formula,'A3DataOnlyCachedValue':cached,'fullCalcOnLoad':calc}
result={'reader':'openpyxl','version':openpyxl.__version__,'checks':[inspect('original.xlsx','=$A$2-$A$1',4),inspect('formula-edited.xlsx','=SUM(A1:A2)',None)],'conclusion':'Both real generated OOXML files opened independently; original formula/cache preserved and edited formula has no stale cache while requesting recalculation.'}
print(json.dumps(result,indent=2))
(Path(__file__).parent/'openpyxl-reader-output.json').write_text(json.dumps(result,indent=2)+'\n')
