from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook, load_workbook
import tempfile, re, sys
root=Path(tempfile.mkdtemp(prefix="xlsx-formula-cache-reader-control-"))
base=root/"base.xlsx"
w=Workbook(); w.active["A1"]="placeholder"; w.save(base)
with ZipFile(base) as zin:
    parts={i.filename:zin.read(i.filename) for i in zin.infolist()}
sheet=parts["xl/worksheets/sheet1.xml"].decode()
sheet=re.sub(r'<c r="A1"[^>]*>.*?</c>', '<c r="A1" TYPE_ATTR><f>1+1</f><v>2</v></c>', sheet)
print("python="+sys.version.split()[0])
import openpyxl
print("openpyxl="+openpyxl.__version__)
for t in ("str","n",None):
    label=t or "omitted"
    attr=f't="{t}"' if t else ""
    parts["xl/worksheets/sheet1.xml"]=sheet.replace("TYPE_ATTR",attr).encode()
    out=root/f"{label}.xlsx"
    with ZipFile(out,"w",ZIP_DEFLATED) as z:
        for name,data in parts.items(): z.writestr(name,data)
    cell=load_workbook(out,data_only=True).active["A1"]
    print(f"{label}: value={cell.value!r}; python_type={type(cell.value).__name__}; openpyxl_data_type={cell.data_type}")
