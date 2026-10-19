import hashlib, json, re, zipfile
from pathlib import Path
from openpyxl import load_workbook
path = Path(__file__).parent / 'interop' / 'table-cell-edited.xlsx'
book = load_workbook(path, data_only=False)
sheet = book['Ventas']
assert 'TablaVentas' in sheet.tables
assert sheet.tables['TablaVentas'].ref == 'A1:B2'
assert sheet['B2'].value == '=SUM(A2:A2)'
with zipfile.ZipFile(path) as archive:
    assert archive.testzip() is None
    xml = archive.read('xl/worksheets/sheet1.xml').decode('utf-8')
    cell = re.search(r'<c\b[^>]*\br="B2"[^>]*>(.*?)</c>', xml).group(1)
    assert '<f>SUM(A2:A2)</f>' in cell
    assert '<v>' not in cell
print(json.dumps({'reader':'openpyxl','version':__import__('openpyxl').__version__,'sheet':'Ventas','table':'TablaVentas','range':sheet.tables['TablaVentas'].ref,'cell':'B2','formula':sheet['B2'].value,'cachedValue':'absent','zipCrc':'PASS','bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()},indent=2))
