import hashlib
import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook

base = Path(__file__).parent / "generated"
expected = {
    "structured-table.xlsx": {"range": "A1:B2", "headers": ["Producto & código", "Cantidad"]},
    "table-resized.xlsx": {"range": "A1:C2", "headers": ["Producto", "Cantidad", "Importe"]},
}
results = []
for filename, target in expected.items():
    path = base / filename
    with zipfile.ZipFile(path) as archive:
        crc_error = archive.testzip()
        table_parts = [name for name in archive.namelist() if name.startswith("xl/tables/table") and name.endswith(".xml")]
    assert crc_error is None, (filename, crc_error)
    assert table_parts == ["xl/tables/table1.xml"], (filename, table_parts)
    workbook = load_workbook(path, data_only=False)
    worksheet = workbook["Ventas"]
    assert list(worksheet.tables.keys()) == ["TablaVentas"]
    table = worksheet.tables["TablaVentas"]
    headers = [column.name for column in table.tableColumns]
    assert table.ref == target["range"], (filename, table.ref)
    assert headers == target["headers"], (filename, headers)
    formula = worksheet["B2"].value if filename == "table-resized.xlsx" else None
    workbook.close()
    cached = None
    if filename == "table-resized.xlsx":
        cached_book = load_workbook(path, data_only=True)
        cached = cached_book["Ventas"]["B2"].value
        cached_book.close()
        assert formula == "=1+2", formula
        assert cached == 3, cached
    results.append({
        "file": filename,
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "zipCrc": "PASS",
        "tableName": "TablaVentas",
        "tableRange": table.ref,
        "headers": headers,
        "formulaB2": formula,
        "cachedB2": cached,
    })
print(json.dumps({
    "reader": "openpyxl 3.1.5",
    "syntheticOnly": True,
    "workbooks": results,
}, ensure_ascii=False, indent=2))
