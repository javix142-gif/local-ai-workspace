import hashlib, json, sys, zipfile
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

root = Path(__file__).parent
expected = {
    'fixture-v1.pptx': {'title': 'Informe visual'},
    'fixture-v2.pptx': {'title': 'Informe revisado'},
}
results = []
for filename, expected_values in expected.items():
    path = root / filename
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, f'{filename}: ZIP CRC failure'
        names = set(archive.namelist())
        required = {
            '[Content_Types].xml',
            'ppt/presentation.xml',
            'ppt/slides/slide1.xml',
            'ppt/slides/_rels/slide1.xml.rels',
            'ppt/slideLayouts/slideLayout1.xml',
            'ppt/slideMasters/slideMaster1.xml',
            'ppt/theme/theme1.xml',
            'ppt/media/image1.png',
        }
        assert required <= names, f'{filename}: missing package parts {sorted(required - names)}'
        import xml.etree.ElementTree as ET
        ET.fromstring(archive.read('ppt/slides/_rels/slide1.xml.rels'))
        rel_xml = archive.read('ppt/slides/_rels/slide1.xml.rels').decode('utf-8')
        assert '../slideLayouts/slideLayout1.xml' in rel_xml
        assert '../media/image1.png' in rel_xml
        image_bytes = archive.read('ppt/media/image1.png')
        image_sha = hashlib.sha256(image_bytes).hexdigest()
    presentation = Presentation(str(path))
    assert len(presentation.slides) == 1
    slide = presentation.slides[0]
    assert slide.slide_layout.name == 'Blank'
    text_values = [shape.text for shape in slide.shapes if shape.has_text_frame]
    assert expected_values['title'] in text_values, (filename, text_values)
    assert 'España y una imagen local.' in text_values, (filename, text_values)
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    picture = pictures[0]
    alt_text = picture._element.nvPicPr.cNvPr.get('descr')
    assert alt_text == 'Mapa & datos', (filename, alt_text)
    assert hashlib.sha256(picture.image.blob).hexdigest() == image_sha
    results.append({
        'file': filename,
        'bytes': path.stat().st_size,
        'sha256': digest,
        'zipCrc': 'PASS',
        'slides': len(presentation.slides),
        'layout': slide.slide_layout.name,
        'masters': len(presentation.slide_masters),
        'themePartPresent': True,
        'title': expected_values['title'],
        'body': 'España y una imagen local.',
        'pictures': len(pictures),
        'altText': alt_text,
        'embeddedImageSha256': image_sha,
    })
assert results[0]['embeddedImageSha256'] == results[1]['embeddedImageSha256']
assert results[0]['sha256'] != results[1]['sha256']
assert results[0]['title'] != results[1]['title']
print(json.dumps({'reader': 'python-pptx', 'version': __import__('pptx').__version__, 'roundtrip': 'PASS', 'versions': results}, ensure_ascii=False, indent=2))
