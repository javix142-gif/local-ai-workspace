import hashlib
import json
import zipfile
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).parent

def inspect(name):
    path = ROOT / name
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None, f"ZIP CRC failure: {name}"
        assert "ppt/theme/theme1.xml" in package.namelist()
        assert "ppt/slideLayouts/slideLayout1.xml" in package.namelist()
        assert "ppt/slideMasters/slideMaster1.xml" in package.namelist()
    presentation = Presentation(path)
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        titles = [shape.text for shape in slide.shapes if shape.name == "Title"]
        bodies = [shape.text for shape in slide.shapes if shape.name == "Body"]
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert len(titles) == 1 and len(bodies) == 1
        assert len(pictures) == 1, (name, index, len(pictures))
        picture = pictures[0]
        descriptions = picture._element.xpath(".//p:cNvPr")
        alt = descriptions[0].get("descr")
        image = picture.image.blob
        image_sha = hashlib.sha256(image).hexdigest()
        image_rels = [rel for rel in slide.part.rels.values() if rel.reltype.endswith("/image")]
        assert len(image_rels) == 1
        assert image_rels[0].target_part.blob == image
        slides.append({
            "slide": index,
            "layout": slide.slide_layout.name,
            "title": titles[0],
            "body": bodies[0],
            "pictureCount": 1,
            "altText": alt,
            "imageBytes": len(image),
            "imageSha256": image_sha,
            "relationshipTarget": image_rels[0].target_part.partname,
        })
    return {
        "file": name,
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "zipCrc": "PASS",
        "slideCount": len(presentation.slides),
        "masterCount": len(presentation.slide_masters),
        "themePartPresent": True,
        "slides": slides,
    }

v1 = inspect("fixture-v1.pptx")
v2 = inspect("fixture-v2.pptx")
assert [s["title"] for s in v1["slides"]] == ["Primera", "Segunda"]
assert [s["title"] for s in v2["slides"]] == ["Primera", "Segunda revisada"]
assert [s["altText"] for s in v1["slides"]] == ["Mapa rojo", "Mapa azul"]
assert [s["altText"] for s in v2["slides"]] == ["Mapa rojo", "Mapa azul"]
assert len({s["imageSha256"] for s in v1["slides"]}) == 2
assert [s["imageSha256"] for s in v1["slides"]] == [s["imageSha256"] for s in v2["slides"]]
result = {
    "reader": "python-pptx",
    "version": __import__("pptx").__version__,
    "result": "PASS_SYNTHETIC_TWO_SLIDE_IMAGE_REVISION_ROUNDTRIP",
    "versions": [v1, v2],
    "assertions": [
        "both PPTX packages reopen with two ordered slides",
        "each slide has one image relationship and one picture",
        "slide-specific alt text is retained",
        "two embedded PNG payloads are distinct",
        "both PNG SHA-256 values are unchanged by the title revision",
        "v1 and v2 have valid ZIP CRC, layout, master and theme parts",
    ],
    "limits": ["No PowerPoint/LibreOffice/Android viewer or device validation."]
}
(ROOT / "independent-reader.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
print(json.dumps(result, ensure_ascii=False, indent=2))
