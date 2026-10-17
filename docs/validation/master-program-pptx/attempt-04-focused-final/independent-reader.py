import hashlib, json, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

path = Path('app/build/pptx-image/fixture-v1.pptx')
ns = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
}
raw = path.read_bytes()
with zipfile.ZipFile(path) as package:
    assert package.testzip() is None
    names = set(package.namelist())
    required = ['ppt/media/image1.png', 'ppt/slides/slideLayout1.xml', 'ppt/slideMasters/slideMaster1.xml', 'ppt/theme/theme1.xml']
    # The generated layout is stored under slideLayouts (plural), matching its relationship target.
    required[1] = 'ppt/slideLayouts/slideLayout1.xml'
    assert all(name in names for name in required), [name for name in required if name not in names]
    types = ET.fromstring(package.read('[Content_Types].xml'))
    assert any(x.attrib.get('Extension') == 'png' and x.attrib.get('ContentType') == 'image/png' for x in types)
    rels = ET.fromstring(package.read('ppt/slides/_rels/slide1.xml.rels'))
    relationships = {x.attrib['Id']: x.attrib for x in rels}
    assert relationships['layout']['Target'] == '../slideLayouts/slideLayout1.xml'
    assert relationships['rIdImage']['Type'].endswith('/image')
    assert relationships['rIdImage']['Target'] == '../media/image1.png'
    slide_xml = ET.fromstring(package.read('ppt/slides/slide1.xml'))
    picture = slide_xml.find('.//p:pic', ns)
    assert picture is not None
    alt = picture.find('./p:nvPicPr/p:cNvPr', ns).attrib['descr']
    assert alt == 'Mapa & datos'
    assert picture.find('./p:blipFill/a:blip', ns).attrib['{%s}embed' % ns['r']] == 'rIdImage'
    ext = picture.find('./p:spPr/a:xfrm/a:ext', ns)
    width, height = int(ext.attrib['cx']), int(ext.attrib['cy'])
    assert abs(width / height - 2.0) < 0.02
    png = package.read('ppt/media/image1.png')
    image = Image.open(path.open('rb')) if False else Image.open(__import__('io').BytesIO(png))
    image.verify()
    assert image.size == (2, 1)

presentation = Presentation(str(path))
assert len(presentation.slides) == 1
slide = presentation.slides[0]
pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
assert len(pictures) == 1
assert slide.slide_layout.name == 'Blank'
assert any('Informe visual' in shape.text for shape in slide.shapes if shape.has_text_frame)
assert any('España y una imagen local.' in shape.text for shape in slide.shapes if shape.has_text_frame)
result = {
    'source_note': 'Artifact generated during attempt-04 before the XML namespace assertion failed; ArtifactRenderer source was unchanged in the subsequent test-harness corrections.',
    'artifact': str(path), 'bytes': len(raw), 'sha256': hashlib.sha256(raw).hexdigest(),
    'zip_crc': 'PASS', 'python_pptx': 'PASS', 'slides': len(presentation.slides),
    'layout': slide.slide_layout.name, 'master_count': len(presentation.slide_masters),
    'theme_part_present': True, 'picture_count': len(pictures), 'image_alt_text': alt,
    'image_bytes': len(png), 'image_sha256': hashlib.sha256(png).hexdigest(),
    'image_dimensions_px': list(image.size), 'shape_aspect_ratio': width / height,
    'slide_size_emu': [presentation.slide_width, presentation.slide_height],
    'title_and_body_text': 'PASS',
}
Path('docs/validation/master-program-pptx/attempt-04-focused-final/independent-reader.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(result, ensure_ascii=False, indent=2))
