import pathlib,json,hashlib,zipfile,importlib.metadata,subprocess
from openpyxl import load_workbook
from docx import Document
from pptx import Presentation
from pypdf import PdfReader
root=pathlib.Path('app/build/artifact-roundtrip');checks=[]
sha=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()
run=json.loads(pathlib.Path('/workspace/master-program-run/listing-final/command.json').read_text())
for f in sorted(root.glob('v*-fixture.*')):
 if f.name.endswith('.json'):continue
 assert f.stat().st_mtime >= run['startEpoch'], f'old artifact output: {f}'
 meta=json.loads(pathlib.Path(str(f)+'.metadata.json').read_text());actual=hashlib.sha256(f.read_bytes()).hexdigest();assert actual==meta['sha256']
 revised=meta['version']==2
 if f.suffix=='.xlsx':
  w=load_workbook(f,data_only=False);assert w.sheetnames==['Hoja1','Hoja2'];assert w['Hoja1']['A2'].value==(7 if revised else 5);assert w['Hoja2']['B14'].value=='Exacto 界';assert w['Hoja1']['A3'].value=='=$A$2-$A$1';assert load_workbook(f,data_only=True)['Hoja1']['A3'].value==4
 elif f.suffix=='.docx':
  d=Document(f);assert d.paragraphs[0].text==('Informe revisado' if revised else 'Informe español');assert d.tables[0].cell(1,0).text=='Álvaro';assert d.paragraphs[1].style.style_id=='Heading1'
 elif f.suffix=='.pptx':
  d=Presentation(f);assert len(d.slides)==2;assert d.slides[0].slide_layout.name=='Blank';assert len(d.slides[0].slide_layout.slide_master.slide_layouts)==1;assert len(d.slide_masters)==1;assert d.slides[0].shapes[0].text==('Informe revisado' if revised else 'Informe español');assert d.slides[1].shapes[1].text=='Álvaro: 42'
 elif f.suffix=='.pdf':
  p=PdfReader(f);assert len(p.pages)==2;assert ('Informe revisado' if revised else 'Informe español') in p.pages[0].extract_text();assert 'Álvaro' in p.pages[1].extract_text()
 elif f.suffix=='.kt':assert f.read_text()==('fun answer() = 43\n' if revised else '// España 界\nfun answer() = 42\n')
 if f.suffix in ['.xlsx','.docx','.pptx']:
  with zipfile.ZipFile(f) as z:assert z.testzip() is None
 checks.append({'file':f.name,'sha256':actual,'bytes':f.stat().st_size,'status':'PASS'})
assert len(checks)==10
result={'sourceSha':sha,'checks':checks,'readers':{p:importlib.metadata.version(p) for p in ['openpyxl','python-docx','python-pptx','pypdf']},'scope':'Installed independent libraries; no physical Android or desktop GUI reader rendering claim'}
pathlib.Path('/workspace/master-program-run/listing-final/independent-reader-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print(json.dumps(result,ensure_ascii=False))
