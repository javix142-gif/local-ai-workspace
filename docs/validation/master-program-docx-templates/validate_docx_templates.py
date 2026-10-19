from __future__ import annotations
import hashlib
import json
import platform
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from docx import Document, __version__ as python_docx_version

ROOT = Path(__file__).parent / "final-focused"
EXPECTED = {
    "report": {"font": "Calibri", "color": "3E746D", "title_pt": 20.0, "heading_pt": 16.0},
    "meeting_minutes": {"font": "Arial", "color": "1D2522", "title_pt": 18.0, "heading_pt": 14.0},
}
EXPECTED_BODY = ["p", "p", "p", "tbl", "sectPr"]
EXPECTED_PARAGRAPHS = ["Título de plantilla", "Acuerdos", "Contenido conservado"]
EXPECTED_TABLE = [["Campo", "Valor"], ["Estado", "Aprobado"]]
results = {"python": platform.python_version(), "python_docx": python_docx_version, "documents": []}
for name, expected in EXPECTED.items():
    path = ROOT / f"{name}.docx"
    data = path.read_bytes()
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"ZIP CRC failure: {name}: {bad_member}")
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise AssertionError(f"duplicate ZIP entry: {name}")
        required_parts = ("word/document.xml", "word/styles.xml", "word/numbering.xml")
        if not all(part in names for part in required_parts):
            raise AssertionError(f"required package part missing: {name}")
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    body = root.find("w:body", ns)
    if body is None:
        raise AssertionError(f"word/document.xml has no body: {name}")
    body_order = [child.tag.rsplit("}", 1)[-1] for child in body]
    if body_order != EXPECTED_BODY:
        raise AssertionError(f"body order mismatch: {name}: {body_order!r}")
    document = Document(path)
    paragraphs = document.paragraphs
    paragraph_text = [p.text for p in paragraphs]
    paragraph_styles = [p.style.name for p in paragraphs]
    table_values = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    title_style = document.styles["Title"]
    heading_style = document.styles["Heading 1"]
    if paragraph_text != EXPECTED_PARAGRAPHS:
        raise AssertionError(f"paragraphs changed: {name}: {paragraph_text!r}")
    if paragraph_styles != ["Title", "Heading 1", "Normal"]:
        raise AssertionError(f"paragraph styles unresolved: {name}: {paragraph_styles!r}")
    if table_values != [EXPECTED_TABLE]:
        raise AssertionError(f"table content changed: {name}: {table_values!r}")
    if title_style.font.name != expected["font"] or heading_style.font.name != expected["font"]:
        raise AssertionError(f"font mismatch: {name}")
    if str(title_style.font.color.rgb) != expected["color"] or str(heading_style.font.color.rgb) != expected["color"]:
        raise AssertionError(f"color mismatch: {name}")
    if title_style.font.size.pt != expected["title_pt"] or heading_style.font.size.pt != expected["heading_pt"]:
        raise AssertionError(f"size mismatch: {name}")
    results["documents"].append({
        "name": path.name,
        "bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "zip_crc": "PASS",
        "required_parts_present": True,
        "direct_body_order": body_order,
        "paragraphs": paragraph_text,
        "paragraph_styles": paragraph_styles,
        "table": table_values[0],
        "title_font_color_size": [title_style.font.name, str(title_style.font.color.rgb), title_style.font.size.pt],
        "heading1_font_color_size": [heading_style.font.name, str(heading_style.font.color.rgb), heading_style.font.size.pt],
    })
print(json.dumps(results, ensure_ascii=False, indent=2))
