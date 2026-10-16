from docx import Document
import hashlib, json, sys

def inspect(path, expected_title):
    doc = Document(path)
    paragraphs = []
    for p in doc.paragraphs:
        numbering = None
        if p._p.pPr is not None and p._p.pPr.numPr is not None:
            num_id = p._p.pPr.numPr.numId
            numbering = num_id.val if num_id is not None else None
        paragraphs.append({
            "text": p.text,
            "style": p.style.name,
            "pageBreakBefore": bool(p.paragraph_format.page_break_before),
            "numberingId": numbering,
        })
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    assert any(p["text"] == expected_title and p["style"] == "Title" for p in paragraphs)
    assert any(p["text"] == "Primera sección" and p["style"] == "Heading 1" for p in paragraphs)
    bullet = next(p for p in paragraphs if p["text"] == "Elemento enumerado")
    assert bullet["numberingId"] == 1
    assert any(p["text"] == "Segunda página" and p["pageBreakBefore"] for p in paragraphs)
    assert tables == [[["Nombre", "Valor"], ["Álvaro", "42"]]]
    with open(path, "rb") as f:
        payload = f.read()
    return {"file": path, "bytes": len(payload), "sha256": hashlib.sha256(payload).hexdigest(),
            "reader": "python-docx 1.2.0", "paragraphs": paragraphs, "tables": tables, "sections": len(doc.sections)}

print(json.dumps([
    inspect(sys.argv[1], "Informe español"),
    inspect(sys.argv[2], "Informe actualizado"),
], ensure_ascii=False, indent=2))
