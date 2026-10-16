from docx import Document
import json, sys
p=sys.argv[1]
doc=Document(p)
paragraphs=[{"text":item.text,"style":item.style.name,"pageBreakBefore":bool(item.paragraph_format.page_break_before)} for item in doc.paragraphs]
tables=[[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
assert any(item["text"]=="Informe español" and item["style"]=="Title" for item in paragraphs)
assert any(item["text"]=="Primera sección" and item["style"]=="Heading 1" for item in paragraphs)
assert any(item["text"]=="Elemento enumerado" and item["style"]=="Normal" for item in paragraphs)
assert any(item["text"]=="Segunda página" and item["pageBreakBefore"] for item in paragraphs)
assert tables == [[ ["Nombre", "Valor"], ["Álvaro", "42"] ]]
print(json.dumps({"reader":"python-docx 1.2.0","file":p,"paragraphs":paragraphs,"tables":tables,"sections":len(doc.sections)},ensure_ascii=False,indent=2))
