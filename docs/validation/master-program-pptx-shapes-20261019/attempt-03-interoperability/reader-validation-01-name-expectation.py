from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation

path = Path('docs/validation/master-program-pptx-shapes-20261019/generated-figures-v1.pptx')
with ZipFile(path) as package:
    assert package.testzip() is None, 'ZIP CRC failure'
presentation = Presentation(str(path))
assert len(presentation.slides) == 1
slide = presentation.slides[0]
expected = {
    'Figure 1': ('RECTANGLE', 'Entrada', 'Entrada & revisión'),
    'Figure 2': ('OVAL', 'Aprobado', 'Resultado aprobado'),
    'Figure 3': ('ROUNDED_RECTANGLE', '', 'Revisión pendiente'),
    'Figure 4': ('TRIANGLE', '', 'Decisión requerida'),
}
for name, (expected_type, expected_text, expected_description) in expected.items():
    shape = next(shape for shape in slide.shapes if shape.name == name)
    assert shape.shape_type == 1, (name, shape.shape_type)
    assert shape.auto_shape_type.name == expected_type, (name, shape.auto_shape_type.name)
    assert shape.text == expected_text, (name, shape.text)
    assert shape._element.nvSpPr.cNvPr.get('descr') == expected_description
    print(f'{name}: type={shape.auto_shape_type.name}; text={shape.text!r}; description={shape._element.nvSpPr.cNvPr.get("descr")!r}')
print(f'PASS python-pptx={__import__("pptx").__version__}; slides={len(presentation.slides)}; zip_crc=PASS; editable_figures={len(expected)}; bytes={path.stat().st_size}')
