#!/usr/bin/env python3
"""Validate the synthetic style-metadata XLSX with ZIP/XML and OpenPyXL."""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument("xlsx", type=Path)
args = parser.parse_args()
path = args.xlsx
main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
with ZipFile(path) as archive:
    bad = archive.testzip()
    assert bad is None, f"bad ZIP CRC: {bad}"
    names = set(archive.namelist())
    required = {
        "[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml",
        "xl/styles.xml", "xl/sharedStrings.xml",
    }
    assert required <= names, f"missing OPC parts: {sorted(required - names)}"
    types = ET.fromstring(archive.read("[Content_Types].xml"))
    content_types = {e.attrib.get("PartName"): e.attrib.get("ContentType") for e in types}
    assert content_types["/xl/workbook.xml"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
    assert content_types["/xl/styles.xml"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"
    assert content_types["/xl/sharedStrings.xml"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
    sheet_xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    raw = {}
    for cell in sheet_xml.findall(f".//{{{main}}}c"):
        formula = cell.find(f"{{{main}}}f")
        value = cell.find(f"{{{main}}}v")
        inline = cell.find(f"{{{main}}}is/{{{main}}}t")
        raw[cell.attrib["r"]] = {
            "type": cell.attrib.get("t", "n"),
            "style": int(cell.attrib.get("s", "0")),
            "formula": formula.text if formula is not None else None,
            "value": value.text if value is not None else None,
            "inline": inline.text if inline is not None else None,
        }
    assert raw["A1"]["value"] == "45292" and raw["A1"]["style"] == 1
    assert raw["B1"]["value"] == "1234.5" and raw["B1"]["style"] == 2
    assert raw["C1"]["type"] == "b" and raw["C1"]["value"] == "1"
    assert raw["D1"]["type"] == "s" and raw["D1"]["value"] == "0"
    assert raw["E1"]["type"] == "inlineStr" and raw["E1"]["inline"] == "Inline label"
    assert raw["F1"]["formula"] == "A1+1" and raw["F1"]["value"] == "45293"

formula_book = openpyxl.load_workbook(path, data_only=False)
cache_book = openpyxl.load_workbook(path, data_only=True)
sheet = formula_book["Ventas"]
cache_sheet = cache_book["Ventas"]
assert formula_book.sheetnames == ["Ventas"]
assert sheet["A1"].value.year == 2024
assert sheet["B1"].value == 1234.5
assert sheet["B1"].number_format == "0.0000;[Red]-0.0000"
assert sheet["C1"].value is True
assert sheet["D1"].value == "Shared label"
assert sheet["E1"].value == "Inline label"
assert sheet["F1"].value == "=A1+1"
assert cache_sheet["F1"].value == 45293
print(json.dumps({
    "result": "PASS",
    "reader": "openpyxl",
    "readerVersion": openpyxl.__version__,
    "zipCrc": "PASS",
    "requiredOpcParts": sorted(required),
    "sheetNames": formula_book.sheetnames,
    "rawCellCount": len(raw),
    "checks": [
        "raw A1 serial/style",
        "raw B1 value/custom number format",
        "boolean cell type",
        "shared string",
        "inline string",
        "formula and stored cache",
        "formula mode does not evaluate",
        "data_only reads cached result",
    ],
    "xlsxBytes": path.stat().st_size,
}, ensure_ascii=False, indent=2))
